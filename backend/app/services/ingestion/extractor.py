"""
Document text extraction service.

Supported formats:
    - PDF
    - TXT
    - Markdown
    - DOCX
    - YAML
    - YML
    - CSV
    - XLSX
    - XLS

The extractor converts every supported file into plain text
so that the downstream chunking + embedding pipeline can
treat all formats consistently.
"""

from io import BytesIO

import codecs

import fitz
import pandas as pd
import yaml
from docx import Document as DocxDocument


# ============================================================
# TEXT DECODING
# ============================================================

def _decode_text_bytes(file_bytes: bytes) -> str:
    """
    Decode raw bytes into text, auto-detecting UTF-16/UTF-8.

    Files saved on Windows are frequently UTF-16 rather than
    UTF-8 — most commonly requirements.txt produced by
    `pip freeze > requirements.txt` in PowerShell, whose
    default output encoding is UTF-16LE. Blindly decoding
    those bytes as UTF-8 doesn't raise an error; it silently
    produces a literal NUL (0x00) character after every ASCII
    character, which Postgres TEXT columns then reject at
    insert time ("PostgreSQL text fields cannot contain NUL
    bytes").

    BOM-based detection is checked first since it's
    unambiguous. Without a BOM, a high proportion of NUL bytes
    in the raw content is a strong signal of unmarked UTF-16.
    """

    if file_bytes.startswith(codecs.BOM_UTF16_LE):
        return file_bytes[2:].decode(
            "utf-16-le",
            errors="replace",
        )

    if file_bytes.startswith(codecs.BOM_UTF16_BE):
        return file_bytes[2:].decode(
            "utf-16-be",
            errors="replace",
        )

    if file_bytes.startswith(codecs.BOM_UTF8):
        return file_bytes.decode(
            "utf-8-sig",
            errors="replace",
        )

    # No BOM. Sample the bytes: UTF-16 text without a BOM is
    # still overwhelmingly full of 0x00 (every other byte, for
    # ASCII-range content), which plain UTF-8 text never is.
    sample = file_bytes[:4096]

    if sample and (sample.count(b"\x00") / len(sample)) > 0.3:

        try:
            return file_bytes.decode(
                "utf-16-le",
                errors="replace",
            )
        except UnicodeDecodeError:
            pass

    return file_bytes.decode(
        "utf-8",
        errors="replace",
    )


def _strip_nul_bytes(text: str) -> str:
    """
    PostgreSQL TEXT/VARCHAR columns cannot store the NUL
    (0x00) character under any encoding. This is a last-resort
    safety net applied to every extractor's output, on top of
    the UTF-16-aware decoding above, so a NUL byte can never
    reach the database regardless of source format or how it
    got there.
    """

    if "\x00" in text:
        return text.replace("\x00", "")

    return text


# ============================================================
# PDF
# ============================================================

def extract_pdf_pages(file_bytes: bytes) -> list[dict]:
    """
    Extract text from a PDF while preserving page numbers.

    Returns:
        [
            {
                "page_number": 1,
                "text": "Text from page 1..."
            },
            ...
        ]
    """

    document = fitz.open(
        stream=file_bytes,
        filetype="pdf",
    )

    pages = []

    try:

        for page_index, page in enumerate(document):

            text = page.get_text()

            if text and text.strip():

                pages.append(
                    {
                        "page_number": page_index + 1,
                        "text": _strip_nul_bytes(
                            text.strip()
                        ),
                    }
                )

    finally:

        document.close()

    return pages


# ============================================================
# TXT
# ============================================================

def extract_txt_text(
    file_bytes: bytes,
) -> str:
    """
    Extract text from a plain text file.
    """

    return _strip_nul_bytes(
        _decode_text_bytes(file_bytes)
    )


# ============================================================
# MARKDOWN
# ============================================================

def extract_md_text(
    file_bytes: bytes,
) -> str:
    """
    Extract Markdown as plain UTF-8 text.

    We intentionally preserve Markdown syntax because
    headings, lists and other structure can be useful
    to the RAG pipeline.
    """

    return _strip_nul_bytes(
        _decode_text_bytes(file_bytes)
    )


def extract_yaml_text(
    file_bytes: bytes,
) -> str:
    """
    Extract YAML content.

    YAML is parsed first to validate that it is valid YAML,
    then converted into readable text.

    If parsing produces a scalar/string, the original
    representation is returned.
    """

    text = _decode_text_bytes(file_bytes)

    if not text.strip():
        return ""

    try:

        data = yaml.safe_load(text)

        if data is None:
            return ""

        if isinstance(data, str):
            return data

        return yaml.safe_dump(
            data,
            sort_keys=False,
            allow_unicode=True,
        )

    except yaml.YAMLError as exc:

        raise ValueError(
            f"Invalid YAML document: {exc}"
        ) from exc


# ============================================================
# CSV
# ============================================================

def extract_csv_text(
    file_bytes: bytes,
) -> str:
    """
    Extract CSV data and convert it into readable text.

    Example:

        name,department,status
        John,Engineering,Active

    becomes:

        name: John | department: Engineering | status: Active
    """

    if not file_bytes:
        return ""

    try:

        dataframe = pd.read_csv(
            BytesIO(file_bytes)
        )

    except Exception as exc:

        raise ValueError(
            f"Unable to read CSV file: {exc}"
        ) from exc

    if dataframe.empty:
        return ""

    rows = []

    for _, row in dataframe.iterrows():

        fields = []

        for column, value in row.items():

            if pd.isna(value):
                continue

            fields.append(
                f"{column}: {value}"
            )

        if fields:
            rows.append(
                " | ".join(fields)
            )

    return "\n".join(rows)


# ============================================================
# EXCEL
# ============================================================

def extract_excel_text(
    file_bytes: bytes,
    file_type: str,
) -> str:
    """
    Extract data from XLSX/XLS files.

    All worksheets are processed.

    The output includes the worksheet name so that
    the source context is preserved.
    """

    if not file_bytes:
        return ""

    try:

        excel_file = pd.ExcelFile(
            BytesIO(file_bytes)
        )

    except Exception as exc:

        raise ValueError(
            f"Unable to read Excel file: {exc}"
        ) from exc

    sections = []

    for sheet_name in excel_file.sheet_names:

        try:

            dataframe = pd.read_excel(
                excel_file,
                sheet_name=sheet_name,
            )

        except Exception as exc:

            raise ValueError(
                f"Unable to read Excel sheet "
                f"'{sheet_name}': {exc}"
            ) from exc

        if dataframe.empty:
            continue

        sections.append(
            f"=== Sheet: {sheet_name} ==="
        )

        for _, row in dataframe.iterrows():

            fields = []

            for column, value in row.items():

                if pd.isna(value):
                    continue

                fields.append(
                    f"{column}: {value}"
                )

            if fields:
                sections.append(
                    " | ".join(fields)
                )

    return "\n".join(sections)


# ============================================================
# DOCX
# ============================================================

def extract_docx_text(
    file_bytes: bytes,
) -> str:
    """
    Extract paragraphs and table content from DOCX.
    """

    try:

        document = DocxDocument(
            BytesIO(file_bytes)
        )

    except Exception as exc:

        raise ValueError(
            f"Unable to read DOCX file: {exc}"
        ) from exc

    sections = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):

        sections.append(
            f"=== Table {table_index} ==="
        )

        for row in table.rows:

            values = []

            for cell in row.cells:

                value = cell.text.strip()

                if value:
                    values.append(value)

            if values:
                sections.append(
                    " | ".join(values)
                )

    return "\n".join(sections)


# ============================================================
# GENERIC TEXT EXTRACTION
# ============================================================

def extract_text(
    file_bytes: bytes,
    file_type: str,
) -> str:
    """
    Generic text extraction function.

    Converts supported document formats into text.
    """

    return _strip_nul_bytes(
        _extract_text_raw(
            file_bytes,
            file_type,
        )
    )


def _extract_text_raw(
    file_bytes: bytes,
    file_type: str,
) -> str:
    """
    Dispatches to the format-specific extractor. Kept separate
    from extract_text() so every branch below is covered by a
    single _strip_nul_bytes() safety net, rather than needing
    the check repeated in each extractor (PDF text, DOCX cell
    text, or an Excel/CSV cell value can in principle also
    surface an embedded NUL byte, not just a raw-text decode).
    """

    file_type = (
        file_type
        .lower()
        .strip()
        .lstrip(".")
    )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if file_type == "pdf":

        pages = extract_pdf_pages(
            file_bytes
        )

        return "\n\n".join(
            page["text"]
            for page in pages
        )

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    if file_type == "txt":

        return extract_txt_text(
            file_bytes
        )

    # --------------------------------------------------------
    # Markdown
    # --------------------------------------------------------

    if file_type in {
        "md",
        "markdown",
    }:

        return extract_md_text(
            file_bytes
        )

    # --------------------------------------------------------
    # YAML
    # --------------------------------------------------------

    if file_type in {
        "yaml",
        "yml",
    }:

        return extract_yaml_text(
            file_bytes
        )

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    if file_type == "csv":

        return extract_csv_text(
            file_bytes
        )

    # --------------------------------------------------------
    # Excel
    # --------------------------------------------------------

    if file_type in {
        "xlsx",
        "xls",
    }:

        return extract_excel_text(
            file_bytes,
            file_type,
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if file_type == "docx":

        return extract_docx_text(
            file_bytes
        )

    # --------------------------------------------------------
    # Unsupported
    # --------------------------------------------------------

    raise ValueError(
        f"Unsupported file type: {file_type}"
    )
